"""Regenerate the optional offline Excel calculator. Requires openpyxl (not a website dependency)."""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.properties import CalcProperties
from openpyxl.utils import get_column_letter
w=Workbook();s=w.active;s.title='日本预估利润表';p=w.create_sheet('参数与逻辑');notes=w.create_sheet('使用说明')
params=[('海运人民币/CBM',800),('国内人民币/CBM',200),('税基系数',.29),('税率',.13),('贴标人民币/箱',4),('清关人民币/整票',1500),('报关人民币/整票',350),('日元兑人民币',.047)]
p.append(['参数','暂定值（可调整）'])
for row in params:p.append(row)
p.append(['固定费分摊','(清关+报关) × 单款CBM / 整票CBM，仅计一次'])
p.append(['税金人民币','申报售价JPY × 数量 × 税基系数 × 税率 × 汇率'])
p.append(['贡献利润','销售额人民币 − 采购成本 − 头程 − 平台佣金及尾程'])
s.append(['日本头程预估贡献利润表 · 一张表只计算一票货']);s.append(['蓝色 A–I 列为输入，绿色 J–V 列为公式。删除示例行输入值后填真实数据；每行数据须完整，否则显示 #N/A。']);s.append(['这是离线计算器，不是网站上传模板。网站上传请下载「头程利润预估」CSV 模板；英国、德国规则尚未提供。'])
headers=['产品名称','MSKU','数量','总CBM','箱数','出厂价RMB/件','申报售价JPY/件','销售价JPY/件','平台佣金及尾程JPY/件','海运RMB','国内RMB','税金JPY','税金RMB','贴标RMB','固定费分摊RMB','头程总RMB','单件头程RMB','采购总RMB','销售额RMB','平台总RMB','贡献利润RMB','贡献利润率']
s.append(headers)
samples=[['示例产品1','EXAMPLE-1',100,1,10,20,2000,3000,800],['示例产品2','EXAMPLE-2',100,3,10,20,2000,3000,800]]
for i,row in enumerate(samples,5):
 for c,value in enumerate(row,1):s.cell(i,c,value)
for r in range(5,25):
 s.cell(r,10,f'=IF(COUNTA(A{r}:I{r})=0,"",IF(OR(COUNT(C{r}:I{r})<>7,C{r}<=0,D{r}<=0,E{r}<=0),NA(),D{r}*\'参数与逻辑\'!$B$2))')
 formulas={11:f'D{r}*\'参数与逻辑\'!$B$3',12:f'G{r}*C{r}*\'参数与逻辑\'!$B$4*\'参数与逻辑\'!$B$5',13:f'L{r}*\'参数与逻辑\'!$B$9',14:f'E{r}*\'参数与逻辑\'!$B$6',15:f'(\'参数与逻辑\'!$B$7+\'参数与逻辑\'!$B$8)*D{r}/SUM($D$5:$D$24)',16:f'J{r}+K{r}+M{r}+N{r}+O{r}',17:f'P{r}/C{r}',18:f'F{r}*C{r}',19:f'H{r}*C{r}*\'参数与逻辑\'!$B$9',20:f'I{r}*C{r}*\'参数与逻辑\'!$B$9',21:f'S{r}-R{r}-P{r}-T{r}',22:f'IF(S{r}=0,"",U{r}/S{r})'}
 for c,formula in formulas.items():s.cell(r,c,f'=IF(J{r}="","",{formula})')
 for c in range(1,23):
  cell=s.cell(r,c);cell.fill=PatternFill('solid',fgColor='EDF5FF' if c<=9 else 'EDF7EF');cell.number_format='0.00' if c>=4 else 'General'
 s.cell(r,22).number_format='0.00%'
s.cell(26,1,'整票合计（人民币，税金 L 列为日元）')
for c in [3,4,5,10,11,12,13,14,15,16,18,19,20,21]:s.cell(26,c,f'=SUM({get_column_letter(c)}5:{get_column_letter(c)}24)')
s.cell(26,22,'=IF(S26=0,"",U26/S26)');s.cell(26,22).number_format='0.00%'
notes.append(['使用说明'])
for text in ['每张表对应一票货，可填 20 款产品。不要将多票费用放入同表。','价格币种：出厂价人民币，申报售价/销售价/平台单件费日元；汇率人民币/日元。','固定清关 1500 + 报关 350 为整票人民币费用，按 CBM 占比分摊。','删除两条示例行 A–I 的输入值后再填真实记录；不要删除公式列。','贡献利润不含广告、仓储、退款等未录入费用，不是最终净利润。','预估参数暂定，不是实际账单。实际加权成本使用网站「实际头程账单」功能。','历史加权公式：(历史基数数量×历史单价 + 新批次实际总额) ÷ 总数量。','历史基数截至日期后才计入新实际账单，避免重复成本。','不支持英国、德国：需要用户提供指定技能正文后实现。','Excel 打开时自动重新计算公式；本文件不会读取或同步网站库存。']:notes.append([text])
for sheet in w:
 sheet.freeze_panes='C5' if sheet==s else 'A2'
 for cell in sheet[4 if sheet==s else 1]:cell.font=Font(color='FFFFFF',bold=True);cell.fill=PatternFill('solid',fgColor='197766');cell.alignment=Alignment(wrap_text=True)
 sheet.row_dimensions[4 if sheet==s else 1].height=34
 for col in sheet.columns:sheet.column_dimensions[col[0].column_letter].width=19 if sheet==s else 44
p.column_dimensions['B'].width=80;notes.column_dimensions['A'].width=110
for row in notes:
 for cell in row:cell.alignment=Alignment(wrap_text=True,vertical='top')
s.auto_filter.ref='A4:V24';w.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True,forceFullCalc=True)
root=Path(__file__).resolve().parents[1];w.save(root/'templates'/'日本头程费用预估利润表.xlsx')
print('Created Excel calculator with editable parameters and 20 formula rows')
